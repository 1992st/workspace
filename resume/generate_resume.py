#!/usr/bin/env python3
"""Generate resume .docx from experience data."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()

# --- Page margins ---
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

# ==================== HELPER ====================
def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = '微软雅黑'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    return h

def add_para(text, bold=False, italic=False, size=Pt(10.5), color=None, space_after=Pt(6)):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = size
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    p.paragraph_format.space_after = space_after
    p.paragraph_format.space_before = Pt(0)
    return p

def add_bullet(text, bold_prefix="", size=Pt(10.5)):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.name = '微软雅黑'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        run.font.size = size
        run = p.add_run(text)
        run.font.name = '微软雅黑'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        run.font.size = size
    else:
        run = p.add_run(text)
        run.font.name = '微软雅黑'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        run.font.size = size
    p.paragraph_format.space_after = Pt(2)
    return p

# ==================== HEADER ====================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('个人简历 — 嵌入式 AI 应用工程师')
run.bold = True
run.font.size = Pt(24)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
run.font.name = '微软雅黑'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('AI Agent 架构 · 嵌入式系统 · LLM 应用开发')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
run.font.name = '微软雅黑'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

# Divider
doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ==================== 个人总结 ====================
add_heading_styled('核心优势', level=2)
add_para('AI + 嵌入式复合型工程师：打通 AI Agent 系统与嵌入式硬件的连接。既有嵌入式底层经验（C/C++、Linux 驱动、交叉编译），又掌握 AI Agent 应用开发（LLM 集成、Agent 工作流、MCP 协议），具备完整的产品化思维和架构设计能力。')

# ==================== 专业技能 ====================
add_heading_styled('专业技能', level=2)

add_para('▎ AI Agent 系统架构', bold=True, size=Pt(11), space_after=Pt(2))
add_bullet('独立构建多 Agent 生态系统：Stone（多 Agent 协调/Win_Stock（股票分析/P100-recordpen（项目管理）')
add_bullet('Agent 工作流、MCP 协议、Function Calling、RAG、Memory 分层管理')
add_bullet('多 Agent 状态监控、记忆继承（MemoryStore + ContextBuilder）、飞书集成')

add_para('▎ 嵌入式系统与音视频', bold=True, size=Pt(11), space_after=Pt(2))
add_bullet('RK3588 平台开发：VI/VP/VO 视频处理流水线、RTSP 多路拼接')
add_bullet('BLE 通信协议：TL7519 GATT 服务设计（20+ Characteristic）')
add_bullet('音频处理器方案：DSP 算法（EQ/DRC/AEC）、I2S/I2C/SPI 驱动')
add_bullet('交叉编译、Buildroot Package 配置、嵌入式 Linux 部署')

add_para('▎ LLM 应用开发', bold=True, size=Pt(11), space_after=Pt(2))
add_bullet('Prompt Engineering、RAG 系统设计、向量数据库')
add_bullet('LLM 替代传统代码评分：市场情绪理解、板块轮动分析、技术指标综合')
add_bullet('AI Agent 记忆系统：Subagent 上下文继承，避免重复错误')

# ==================== 项目经验 ====================
add_heading_styled('项目经验', level=2)

# Project 1
add_para('1. 嵌入式视频处理引擎 (RK3588)', bold=True, size=Pt(11), space_after=Pt(2))
add_para('2025 — 至今', italic=True, size=Pt(9), color=RGBColor(0x88, 0x88, 0x88), space_after=Pt(2))
add_bullet('基于 RK3588 平台实现 RTSP 多路实时拼接，开发 VI/VP/VO 完整流水线')
add_bullet('搭建 Mac → Linux 服务器 → RK3588 三端交叉编译部署流程')
add_bullet('目标：用自研代码替代原有 SDK，实现自主可控的视频处理能力')

# Project 2
add_para('2. 智能录音笔 BLE 协议栈 (P100)', bold=True, size=Pt(11), space_after=Pt(2))
add_para('2025 — 至今', italic=True, size=Pt(9), color=RGBColor(0x88, 0x88, 0x88), space_after=Pt(2))
add_bullet('TL7519 + TL9118 双芯片架构，负责 BLE 通信协议设计和文档输出')
add_bullet('实现 9 个自定义 GATT 服务和 20+ Characteristic（时间同步、文件传输、设备控制）')
add_bullet('完成 C100 vs P100 BLE 协议对比分析，输出完整设计文档')

# Project 3
add_para('3. AI Agent 多 Agent 协作平台 (Stone)', bold=True, size=Pt(11), space_after=Pt(2))
add_para('2025 — 至今', italic=True, size=Pt(9), color=RGBColor(0x88, 0x88, 0x88), space_after=Pt(2))
add_bullet('设计并实现 Business Partner Agent，协调管理 5+ 专业 Agent')
add_bullet('状态监控系统（自适应监控周期）、Agent 记忆继承、Gateway HTTP API 唤醒机制')
add_bullet('飞书群集成：监控报告推送 + 双向沟通')

# Project 4
add_para('4. A 股智能投资分析系统 (Win_Stock)', bold=True, size=Pt(11), space_after=Pt(2))
add_para('2025 — 至今', italic=True, size=Pt(9), color=RGBColor(0x88, 0x88, 0x88), space_after=Pt(2))
add_bullet('数据管道（AkShare， LLM 四阶段分析法（大盘→板块→新闻→个股）')
add_bullet('LLM 综合分析替代传统评分系统，提升决策灵活性')
add_bullet('每日复盘机制，积累个股股性理解，15+ 自选股持续跟踪')

# Project 5
add_para('5. 音频处理器产品全流程管理', bold=True, size=Pt(11), space_after=Pt(2))
add_para('2026', italic=True, size=Pt(9), color=RGBColor(0x88, 0x88, 0x88), space_after=Pt(2))
add_bullet('芯片方案调研对比、技术方案设计、嵌入式固件开发')
add_bullet('覆盖完整产品生命周期：市场调研→方案设计→开发实现→量产跟踪')

# ==================== 技术栈 ====================
add_heading_styled('技术栈', level=2)

table = doc.add_table(rows=5, cols=2, style='Light Grid Accent 1')
table.alignment = WD_TABLE_ALIGNMENT.CENTER
data = [
    ('编程语言', 'C/C++, Python, Go, Node.js, Shell Script, JavaScript'),
    ('嵌入式', 'RK3588, TL7519, Buildroot, Linux 驱动, BLE, I2S/I2C/SPI'),
    ('AI/LLM', 'Agent 框架, MCP 协议, LangChain, RAG, Prompt Engineering, 向量数据库'),
    ('工具平台', 'Git, Docker, SSH/rsync/adb, Feishu API, AkShare, SQLite'),
    ('音视频', 'RTSP/RTMP, FFmpeg, DSP 算法 (EQ/DRC/AEC)'),
]
for i, (key, val) in enumerate(data):
    row = table.rows[i]
    row.cells[0].text = key
    row.cells[1].text = val
    for cell in row.cells:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.size = Pt(10)
                run.font.name = '微软雅黑'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

# ==================== Footer ====================
doc.add_paragraph()
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('— 基于实际项目经验整理 · 2026-05 —')
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
run.font.name = '微软雅黑'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

output_path = '/Volumes/zhangstExtern/openclaw/workspace/stone/resume/简历_嵌入式AI应用工程师.docx'
doc.save(output_path)
print(f'✅ 简历已保存: {output_path}')
