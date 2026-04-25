from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

# Slide dimensions (16:9)
W = Inches(13.333)
H = Inches(7.5)

# Background colors
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(245, 245, 247)

slides = [
    ("01-cover.png", BLACK),
    ("02-overview.png", BLACK),
    ("03-architecture.png", BLACK),
    ("04-prompt-overview.png", BLACK),
    ("05-inject-assemble.png", BLACK),
    ("06-skills.png", BLACK),
    ("07-memory.png", BLACK),
    ("08-optimizations.png", BLACK),
    ("09-p100-case.png", BLACK),
    ("10-summary.png", WHITE),
]

prs = Presentation()
prs.slide_width = W
prs.slide_height = H

import os

for img_name, bg_color in slides:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Set background
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color

    img_path = os.path.join(os.path.dirname(__file__), img_name)
    if os.path.exists(img_path):
        # Center image if aspect ratio differs slightly, but since both are 16:9, fill entire slide
        slide.shapes.add_picture(img_path, Inches(0), Inches(0), width=W, height=H)
    else:
        print(f"Warning: {img_name} not found, left blank.")

output = os.path.join(os.path.dirname(__file__), "OpenClaw_PPT_Final.pptx")
prs.save(output)
print(f"Saved: {output}")

# Verify
size = os.path.getsize(output)
print(f"File size: {size / 1024:.1f} KB")
