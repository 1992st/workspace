from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_textbox(slide, left, top, width, height, text, font_size, font_color, bold=False, font_name="SF Pro Display"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = font_color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = PP_ALIGN.LEFT


def add_code_block(slide, left, top, width, height, text, font_size=14):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(20, 20, 25)
    box.line.color.rgb = RGBColor(50, 50, 60)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = RGBColor(200, 200, 210)
    p.font.name = "SF Mono"
    p.alignment = PP_ALIGN.LEFT


# Page 1: Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0, 0, 0)
add_textbox(slide, 1.5, 2.2, 10, 1.5, "OpenClaw 功能拆解", 72, RGBColor(255,255,255))
add_textbox(slide, 1.5, 3.5, 10, 0.8, "从 Prompt 工程到 Agent 上下文管理", 32, RGBColor(161,161,166))
add_textbox(slide, 5.2, 6.5, 3, 0.5, "Engineering Deep Dive", 16, RGBColor(110,110,115))

# Page 2: Overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0,0,0)
add_textbox(slide, 0.8, 0.6, 10, 1, "OpenClaw 解决什么问题？", 52, RGBColor(255,255,255))
add_textbox(slide, 0.8, 1.4, 12, 0.6, "在 LLM 无状态的前提下，搭建一条可控制的能力链", 24, RGBColor(161,161,166))
add_textbox(slide, 1.0, 3.0, 4, 2, "LLM Stateless\n（断裂链条示意，请导入截图）", 20, RGBColor(136,136,136))
add_textbox(slide, 5.0, 4.0, 3.5, 0.5, "OpenClaw Framework", 20, RGBColor(96,165,250))
add_textbox(slide, 9.0, 3.0, 4, 2.5, "Inject    Assemble    Extend    Persist\n（四个齿轮示意，请导入截图）", 22, RGBColor(200,200,200))

# Page 3: Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0,0,0)
add_textbox(slide, 0.8, 0.6, 12, 1, "五层数据流", 52, RGBColor(255,255,255))
steps = [
    ("1. Gateway", "agentId + sessionKey"),
    ("2. Session", "history.jsonl"),
    ("3. Bootstrap", "contextFiles[]"),
    ("4. System Prompt", "prompt string"),
    ("5. Runtime", "LLM request"),
]
x = 0.6
for title, subtitle in steps:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.8), Inches(2.2), Inches(1.6))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(10,10,12)
    shape.line.color.rgb = RGBColor(68,68,68)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(22)
    p.font.color.rgb = RGBColor(255,255,255)
    p.font.bold = True
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(136,136,136)
    p2.font.name = "SF Mono"
    x += 2.5

# Page 4: Prompt Overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0,0,0)
add_textbox(slide, 0.8, 0.5, 12, 1, "Prompt 总览", 48, RGBColor(255,255,255))
add_textbox(slide, 0.8, 1.2, 12, 0.5, "System Prompt 由三部分拼成", 22, RGBColor(161,161,166))
add_textbox(slide, 0.8, 2.5, 3.8, 3.5, "硬编码规则 ~35%\nSafety\nExecution Bias\nCron\nCLI", 20, RGBColor(200,200,255))
add_textbox(slide, 4.8, 2.5, 3.8, 3.5, "工作区上下文 ~40%\nAGENTS.md\nSOUL.md\nTOOLS.md\nMEMORY.md", 20, RGBColor(255,200,170))
add_textbox(slide, 8.8, 2.5, 3.8, 3.5, "动态环境 ~25%\nWorkspace\nRuntime\nSandbox", 20, RGBColor(220,180,255))
add_textbox(slide, 5.5, 6.2, 3, 0.4, "buildAgentSystemPrompt()", 16, RGBColor(136,136,136))
add_textbox(slide, 6.0, 6.7, 1.5, 0.4, "LLM", 20, RGBColor(255,255,255))

# Page 5: Inject & Assemble
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0,0,0)
add_textbox(slide, 0.6, 0.4, 10, 1, "注入与拼装", 44, RGBColor(255,255,255))
add_textbox(slide, 0.6, 1.0, 12, 0.5, "通过状态标记控制注不注，通过三层结构控制怎么拼", 20, RGBColor(161,161,166))
add_code_block(slide, 0.6, 1.8, 6.5, 1.2, "customType: openclaw:bootstrap-context:full -> Skip\ntype: message (assistant)\ntype: compaction -> Re-inject", 16)
add_textbox(slide, 0.6, 3.1, 4, 0.3, "hasCompletedBootstrapTurn()", 14, RGBColor(102,102,102))
layer_colors = [RGBColor(30,64,175), RGBColor(194,65,12), RGBColor(126,34,206)]
layer_texts = ["硬编码规则\nSafety / Execution Bias", "Project Context\nAGENTS.md / SOUL.md / MEMORY.md", "动态环境\nWorkspace / Runtime"]
x = 0.6
for color, text in zip(layer_colors, layer_texts):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(4.0), Inches(3.8), Inches(1.8))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(255,255,255)
    x += 4.1
add_textbox(slide, 12.5, 4.7, 1, 0.5, "LLM", 22, RGBColor(255,255,255))

# Page 6: Skills
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0,0,0)
add_textbox(slide, 0.6, 0.4, 10, 1, "Skills — 可插拔的能力包", 46, RGBColor(255,255,255))
add_textbox(slide, 0.6, 1.1, 12, 0.4, "匹配 skill -> 必须 read SKILL.md -> 严格执行", 20, RGBColor(161,161,166))
add_textbox(slide, 0.6, 2.0, 3, 3, "docx\nweather\ngithub\ndevice-build-debug（高亮）\npptx\n\n（抽屉柜示意）", 18, RGBColor(221,221,221))
add_textbox(slide, 4.0, 2.5, 3, 2, "SKILL.md\n1. SSH login remote\n2. Compile via make\n3. rsync firmware\n4. adb push & verify", 16, RGBColor(245,158,11))
add_textbox(slide, 7.5, 2.8, 4, 0.8, "Read  ->  Follow  ->  Execute", 20, RGBColor(255,255,255))
add_textbox(slide, 10.5, 2.0, 2.5, 1, "skillsSnapshot\n父传子\n跳过扫描", 14, RGBColor(96,165,250))
add_textbox(slide, 10.0, 5.0, 3, 1.5, "6 级来源优先级\nworkspace/skills（最高）", 16, RGBColor(245,158,11))

# Page 7: Memory
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0,0,0)
add_textbox(slide, 0.6, 0.4, 10, 1, "Memory — 可替换的记忆后端", 46, RGBColor(255,255,255))
add_textbox(slide, 0.6, 1.1, 10, 0.4, "registerMemoryCapability()", 20, RGBColor(161,161,166))
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(3.0), Inches(3.0), Inches(2.2))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(15,10,20)
shape.line.color.rgb = RGBColor(167,139,250)
tf = shape.text_frame
p = tf.paragraphs[0]
p.text = "Memory Plugin Slot"
p.font.size = Pt(22)
p.font.color.rgb = RGBColor(167,139,250)
add_textbox(slide, 4.2, 2.6, 1.8, 1, "落盘\nFiles", 18, RGBColor(96,165,250))
add_textbox(slide, 5.8, 3.8, 2.2, 0.8, "分块\nchunk · hash", 18, RGBColor(245,158,11))
add_textbox(slide, 4.2, 5.0, 1.8, 1, "检索\nVector DB", 18, RGBColor(167,139,250))
add_textbox(slide, 9.5, 3.5, 2.5, 1.5, "(  search  )\n大眼睛示意", 18, RGBColor(255,255,255))
add_textbox(slide, 5.0, 6.5, 1.8, 0.4, "builtin 向量索引", 16, RGBColor(96,165,250))
add_textbox(slide, 7.2, 6.5, 1.5, 0.4, "qmd 外部后端", 16, RGBColor(167,139,250))

# Page 8: Optimizations
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0,0,0)
add_textbox(slide, 0.8, 0.5, 12, 1, "效率与主动性", 48, RGBColor(255,255,255))
add_textbox(slide, 0.8, 1.2, 12, 0.5, "框架在后台默默做的三件事", 22, RGBColor(161,161,166))
add_textbox(slide, 5.8, 2.2, 2.5, 1.2, "Auto-Compaction\n80 轮 -> 3 条 summary", 20, RGBColor(245,158,11))
add_textbox(slide, 2.2, 5.0, 2.2, 1.0, "Snapshot\n子代理秒启动", 20, RGBColor(96,165,250))
add_textbox(slide, 9.0, 5.0, 2.2, 1.0, "Heartbeat\n30 分钟自动巡检", 20, RGBColor(52,211,153))
add_textbox(slide, 5.8, 4.3, 2.5, 0.5, "Invisible to user", 16, RGBColor(85,85,85))

# Page 9: P100 Case
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0,0,0)
add_textbox(slide, 0.5, 0.4, 10, 0.9, "P100-recordpen 编译 + 烧录 + 调试", 38, RGBColor(255,255,255))
add_textbox(slide, 0.5, 1.1, 10, 0.4, "一个请求激活四个子系统", 18, RGBColor(161,161,166))
shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.5), Inches(2.6), Inches(2.4), Inches(2.4))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(10,10,12)
shape.line.color.rgb = RGBColor(255,255,255)
tf = shape.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "帮我编译并\n烧录到设备上\n调试固件"
p.font.size = Pt(16)
p.font.color.rgb = RGBColor(255,255,255)
p.alignment = PP_ALIGN.CENTER
add_textbox(slide, 2.0, 1.8, 2.2, 1.0, "1. Bootstrap\n加载 5 个核心 md", 15, RGBColor(96,165,250))
add_textbox(slide, 9.0, 1.8, 2.2, 1.0, "2. Skills\n命中 device-build-debug", 15, RGBColor(245,158,11))
add_textbox(slide, 1.2, 5.0, 2.2, 1.0, "3. Memory\n检索 IP / 编译参数", 15, RGBColor(167,139,250))
add_textbox(slide, 9.8, 5.0, 2.4, 1.0, "4. System Prompt\n拼 SSH + Sandbox", 15, RGBColor(52,211,153))
add_textbox(slide, 4.5, 6.2, 4.5, 0.8, "SSH (Compile)  ->  rsync (Flash)  ->  adb (Debug)", 18, RGBColor(200,200,200))
add_code_block(slide, 0.5, 5.4, 3.5, 0.8, "device-build-debug", 14)

# Page 10: Summary
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(245,245,247)
add_textbox(slide, 5.2, 1.6, 3, 1.2, "Inject", 90, RGBColor(0,0,0))
add_textbox(slide, 5.0, 2.8, 3.4, 1.2, "Assemble", 90, RGBColor(0,0,0))
add_textbox(slide, 5.3, 4.0, 3, 1.2, "Extend", 90, RGBColor(0,0,0))
add_textbox(slide, 5.2, 5.2, 3, 1.2, "Persist", 90, RGBColor(0,0,0))
add_textbox(slide, 5.0, 6.8, 3.5, 0.4, "Engineering, not magic.", 18, RGBColor(136,136,136))

output_path = '/Volumes/zhangstExtern/openclaw/workspace/agent-radar-desk/dev-topics/_migrated/openclaw-ppt-design/OpenClaw_PPT_Template.pptx'
prs.save(output_path)
print(f"Saved: {output_path}")
